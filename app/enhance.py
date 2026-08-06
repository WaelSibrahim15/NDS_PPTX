"""NDS — Enhance mode: upgrade an existing, well-designed deck in place.

Keeps the uploaded design untouched except for additive improvements:
  * AI-generated imagery (OpenAI Images API) inserted into genuinely EMPTY
    regions of a slide, so nothing ever overlaps the original content;
  * narration + movement are handled by the normal narrate pipeline
    (pptx_builder.narrate_existing_pptx with animate=True).

Empty-region detection: every shape's bounding box is rasterised onto a grid
and the largest empty rectangle is found; an image is only added when that
rectangle is big enough to look intentional.
"""
import base64
import io
import json
import urllib.request
from pathlib import Path
from typing import Callable, List, Optional

from pptx import Presentation

GRID_W, GRID_H = 64, 36
MAX_IMAGES = 6          # cost guard per deck
PAD_FRAC = 0.012        # padding around occupied boxes, fraction of slide width


class EnhanceError(RuntimeError):
    pass


# ------------------------------------------------------------- slide analysis

def _occupied_grid(slide, sw: int, sh: int):
    """Rasterise shape bounding boxes onto a GRID_W x GRID_H boolean grid."""
    grid = [[False] * GRID_W for _ in range(GRID_H)]
    pad = int(sw * PAD_FRAC)
    for sp in slide.shapes:
        try:
            l, t, w, h = sp.left, sp.top, sp.width, sp.height
        except Exception:
            continue
        if l is None or t is None or w is None or h is None:
            continue
        x0 = max(0, int((l - pad) * GRID_W / sw))
        x1 = min(GRID_W - 1, int((l + w + pad) * GRID_W / sw))
        y0 = max(0, int((t - pad) * GRID_H / sh))
        y1 = min(GRID_H - 1, int((t + h + pad) * GRID_H / sh))
        for y in range(y0, y1 + 1):
            for x in range(x0, x1 + 1):
                grid[y][x] = True
    return grid


def _largest_empty_rect(grid):
    """Classic maximal-rectangle-in-binary-matrix (histogram method).
    Returns (x0, y0, x1, y1) in grid cells, or None."""
    best, best_area = None, 0
    heights = [0] * GRID_W
    for y in range(GRID_H):
        for x in range(GRID_W):
            heights[x] = heights[x] + 1 if not grid[y][x] else 0
        stack = []
        for x in range(GRID_W + 1):
            h = heights[x] if x < GRID_W else 0
            start = x
            while stack and stack[-1][1] >= h:
                s, sh_ = stack.pop()
                area = sh_ * (x - s)
                if area > best_area:
                    best_area = area
                    best = (s, y - sh_ + 1, x - 1, y)
                start = s
            stack.append((start, h))
    return best


def analyze_slides(pptx_path: Path) -> List[dict]:
    """Per slide: does it already have a photo, and where is the largest empty
    rectangle (EMU) big enough to host an added image?"""
    prs = Presentation(str(pptx_path))
    sw, sh = prs.slide_width, prs.slide_height
    out = []
    for slide in prs.slides:
        has_image = any(sp.shape_type is not None and "PICTURE" in str(sp.shape_type)
                        for sp in slide.shapes)
        rect_cells = _largest_empty_rect(_occupied_grid(slide, sw, sh))
        rect = None
        if rect_cells:
            x0, y0, x1, y1 = rect_cells
            l = int(x0 * sw / GRID_W)
            t = int(y0 * sh / GRID_H)
            w = int((x1 - x0 + 1) * sw / GRID_W)
            h = int((y1 - y0 + 1) * sh / GRID_H)
            # an added image must be substantial, not a sliver
            if w >= 0.20 * sw and h >= 0.25 * sh:
                rect = {"left": l, "top": t, "width": w, "height": h}
        out.append({"has_image": has_image, "rect": rect})
    return out


# ------------------------------------------------------------ image generation

def _openai_image(api_key: str, prompt: str, size: str) -> bytes:
    """One image via OpenAI; gpt-image-1 first, dall-e-3 as fallback."""
    def call(model, extra):
        payload = {"model": model, "prompt": prompt, "size": size, "n": 1, **extra}
        req = urllib.request.Request(
            "https://api.openai.com/v1/images/generations",
            data=json.dumps(payload).encode(),
            headers={"Authorization": f"Bearer {api_key}",
                     "Content-Type": "application/json"},
        )
        with urllib.request.urlopen(req, timeout=180) as resp:
            data = json.loads(resp.read())
        item = data["data"][0]
        if item.get("b64_json"):
            return base64.b64decode(item["b64_json"])
        with urllib.request.urlopen(item["url"], timeout=120) as r:
            return r.read()

    try:
        return call("gpt-image-1", {"quality": "medium"})
    except Exception:
        d3_size = {"1024x1024": "1024x1024", "1536x1024": "1792x1024",
                   "1024x1536": "1024x1792"}.get(size, "1024x1024")
        try:
            return call("dall-e-3", {"size": d3_size, "response_format": "b64_json"})
        except Exception as exc:
            raise EnhanceError(f"Image generation failed: {exc}")


def _pick_size(rect) -> str:
    ar = rect["width"] / max(1, rect["height"])
    if ar >= 1.3:
        return "1536x1024"
    if ar <= 0.77:
        return "1024x1536"
    return "1024x1024"


STYLE_SUFFIX = (" Professional corporate photography, clean modern aesthetic, cool "
                "blue tones, high quality, realistic. No text, no words, no logos, "
                "no watermarks, no charts.")


def add_images(
    src: Path,
    out: Path,
    prompts: List[Optional[str]],
    rects: List[Optional[dict]],
    api_key: str,
    progress: Callable[[int, str], None] = lambda pct, msg: None,
) -> int:
    """Insert one generated image per slide that has a prompt AND a big-enough
    empty region. Contain-fit inside the region, centered. Returns count."""
    prs = Presentation(str(src))
    added = 0
    todo = [(i, p) for i, p in enumerate(prompts) if p and i < len(rects) and rects[i]]
    todo = todo[:MAX_IMAGES]
    for k, (idx, prompt) in enumerate(todo):
        rect = rects[idx]
        progress(int(100 * k / max(1, len(todo))),
                 f"Creating image {k + 1} of {len(todo)}…")
        png = _openai_image(api_key, prompt + STYLE_SUFFIX, _pick_size(rect))
        img_w, img_h = [int(x) for x in _pick_size(rect).split("x")]
        # contain-fit into the empty rectangle
        scale = min(rect["width"] / img_w, rect["height"] / img_h)
        w, h = int(img_w * scale), int(img_h * scale)
        left = rect["left"] + (rect["width"] - w) // 2
        top = rect["top"] + (rect["height"] - h) // 2
        slide = list(prs.slides)[idx]
        slide.shapes.add_picture(io.BytesIO(png), left, top, w, h)
        added += 1
    if added:
        prs.save(str(out))
    return added
