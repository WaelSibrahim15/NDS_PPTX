"""NDS — pull plain text out of the uploaded source document."""
from pathlib import Path

MAX_CHARS = 400_000  # keep well inside the model context window


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        text = path.read_text(errors="replace")
    elif suffix == ".docx":
        text = _from_docx(path)
    elif suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix == ".pptx":
        text = _from_pptx(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix} (use .docx, .pdf, .pptx, .txt or .md)")

    text = text.strip()
    if not text:
        raise ValueError("No readable text found in the uploaded file.")
    if len(text) > MAX_CHARS:
        raise ValueError(
            f"Document is too large ({len(text):,} characters; limit {MAX_CHARS:,}). "
            "Split it or upload a shorter extract."
        )
    return text


def _from_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_slides(path: Path) -> list:
    """Per-slide text + speaker notes from an existing deck (for narrate mode)."""
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"index": i, "texts": texts, "notes": notes})
    if not slides:
        raise ValueError("That PowerPoint has no slides.")
    return slides


def _from_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(f"[Speaker notes] {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(parts)
