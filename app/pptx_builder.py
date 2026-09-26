"""NDS — assemble the narrated .pptx from a deck plan + per-slide audio files.

Slide design lives in design.py, coloured by the chosen template (themes.py);
this module turns those layouts into PowerPoint shapes and adds narration.
"""
from pathlib import Path
from typing import List, Optional

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import design
from .models import DeckPlan
from .tts import audio_duration_seconds

# python-pptx (≤1.0.x) crashes when adding media to a deck that already contains
# media/images loaded as plain Parts (no .sha1). Same guard the library already
# uses for images (pptx/package.py _ImageParts) applied to media parts.
from pptx.package import _MediaParts


def _find_media_by_sha1_safe(self, sha1):
    for media_part in self:
        if not hasattr(media_part, "sha1"):
            continue
        if media_part.sha1 == sha1:
            return media_part
    return None


_MediaParts._find_by_sha1 = _find_media_by_sha1_safe

SLIDE_W = Inches(design.W)
SLIDE_H = Inches(design.H)


def build_deck(
    plan: DeckPlan,
    audio_files: List[Optional[Path]],
    out_path: Path,
    template="niq",
) -> Path:
    """template: a template id or a theme dict resolved by themes."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for idx, spec in enumerate(plan.slides):
        slide = prs.slides.add_slide(blank)
        _draw_scene(slide, design.layout_slide(plan, idx, template))

        # narration into speaker notes so a human presenter can reuse the deck
        slide.notes_slide.notes_text_frame.text = spec.narration

        audio = audio_files[idx] if idx < len(audio_files) else None
        if audio is not None and Path(audio).exists():
            duration = audio_duration_seconds(Path(audio))
            _embed_autoplay_audio(slide, Path(audio))
            _set_auto_advance(slide, duration + 1.0)
        # Silent drafts: no auto-advance — presenter (or reviewer) advances manually.

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


# ------------------------------------------------------------ scene → shapes

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _no_shadow_solid(shp, fill: Optional[str]):
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
    else:
        shp.fill.background()
    shp.shadow.inherit = False


def _draw_scene(slide, scene: "design.Scene"):
    """Add every element of the scene to the slide."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(scene.background)
    for el in scene.elements:
        _drop_theme_style(_draw_element(slide, el, scene.fonts))


def _drop_theme_style(shp):
    """Remove the theme style reference python-pptx adds to shapes and lines, so no
    theme shadow or outline shows up: the design is flat colour only."""
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)


def _draw_element(slide, el, fonts: dict):
    shapes = slide.shapes
    if isinstance(el, design.Rect):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if el.radius else MSO_SHAPE.RECTANGLE
        shp = shapes.add_shape(kind, Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h))
        _no_shadow_solid(shp, el.fill)
        shp.line.fill.background()
        if el.radius:
            shp.adjustments[0] = min(0.5, el.radius / min(el.w, el.h))
        return shp
    if isinstance(el, design.Oval):
        d = Inches(el.r * 2)
        shp = shapes.add_shape(MSO_SHAPE.OVAL, Inches(el.cx - el.r), Inches(el.cy - el.r), d, d)
        _no_shadow_solid(shp, el.fill)
        if el.line:
            shp.line.color.rgb = _rgb(el.line)
            shp.line.width = Pt(el.line_w)
        else:
            shp.line.fill.background()
        return shp
    if isinstance(el, design.Arc):
        d = Inches(el.r * 2)
        shp = shapes.add_shape(MSO_SHAPE.ARC, Inches(el.cx - el.r), Inches(el.cy - el.r), d, d)
        # python-pptx scales raw adjustment values by 1/100000; angles are in 60000ths of a degree
        shp.adjustments[0] = (el.start % 360) * 0.6
        shp.adjustments[1] = (el.end % 360) * 0.6
        _no_shadow_solid(shp, None)
        shp.line.color.rgb = _rgb(el.color)
        shp.line.width = Pt(el.line_w)
        return shp
    if isinstance(el, design.Line):
        shp = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(el.x1), Inches(el.y1),
                                   Inches(el.x2), Inches(el.y2))
        shp.line.color.rgb = _rgb(el.color)
        shp.line.width = Pt(el.line_w)
        return shp
    if isinstance(el, design.Image):
        return shapes.add_picture(str(el.path), Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h))
    if isinstance(el, design.Text):
        box = shapes.add_textbox(Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = _ANCHOR[el.anchor]
        for i, para in enumerate(el.paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _ALIGN[para.align]
            p.line_spacing = para.line_spacing
            if para.space_after:
                p.space_after = Pt(para.space_after)
            for run in para.runs:
                r = p.add_run()
                r.text = run.text
                r.font.name = fonts.get(run.font) or "Arial"
                r.font.size = Pt(run.size)
                r.font.bold = run.bold
                r.font.italic = run.italic
                r.font.color.rgb = _rgb(run.color)
        return box
    raise TypeError(f"Unknown design element: {el!r}")


def narrate_existing_pptx(
    src_path: Path,
    narrations: List[str],
    audio_files: List[Optional[Path]],
    out_path: Path,
) -> Path:
    """Embed narration audio into a COPY of an existing deck, leaving its design
    untouched: per-slide auto-play audio, auto-advance, narration in the notes."""
    prs = Presentation(str(src_path))
    for idx, slide in enumerate(prs.slides):
        if idx < len(narrations) and narrations[idx].strip():
            slide.notes_slide.notes_text_frame.text = narrations[idx]
        audio = audio_files[idx] if idx < len(audio_files) else None
        if audio is not None and Path(audio).exists():
            duration = audio_duration_seconds(Path(audio))
            _embed_autoplay_audio(slide, Path(audio))
            _set_auto_advance(slide, duration + 1.0)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


# ------------------------------------------------------- audio embed + timing

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"

# Slide timing tree: the media playback lives inside ONE group time-node that
# begins when the slide appears (onBegin of the main sequence), so it fires
# automatically under an auto-advancing show, with no click needed.
_TIMING_HEAD = f"""<p:timing xmlns:p="{_P}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <p:tnLst>
  <p:par>
   <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
    <p:childTnLst>
     <p:seq concurrent="1" nextAc="seek">
      <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
       <p:childTnLst>
        <p:par>
         <p:cTn id="3" fill="hold">
          <p:stCondLst>
           <p:cond delay="indefinite"/>
           <p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>
          </p:stCondLst>
          <p:childTnLst>"""

# Media playback child (starts immediately, delay 0). ids {i0}-{i2}.
_TIMING_MEDIA = """           <p:par>
            <p:cTn id="{i0}" fill="hold">
             <p:stCondLst><p:cond delay="0"/></p:stCondLst>
             <p:childTnLst>
              <p:par>
               <p:cTn id="{i1}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                <p:childTnLst>
                 <p:cmd type="call" cmd="playFrom(0.0)">
                  <p:cBhvr>
                   <p:cTn id="{i2}" dur="{dur}" fill="hold"/>
                   <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                  </p:cBhvr>
                 </p:cmd>
                </p:childTnLst>
               </p:cTn>
              </p:par>
             </p:childTnLst>
            </p:cTn>
           </p:par>"""

_TIMING_TAIL = """          </p:childTnLst>
         </p:cTn>
        </p:par>
       </p:childTnLst>
       <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
       <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
      </p:cTn>
     </p:seq>
    </p:childTnLst>
   </p:cTn>
  </p:par>
 </p:tnLst>
</p:timing>"""

_MIME_BY_SUFFIX = {".mp3": "audio/mpeg", ".m4a": "audio/mp4", ".wav": "audio/wav"}


def _build_timing_xml(spid: int, dur_ms: int) -> str:
    """Assemble the slide timing tree that auto-plays the narration shape."""
    # ids 1-3 are the header (tmRoot, mainSeq, group). Media child uses 4-6.
    pars = _TIMING_MEDIA.format(i0=4, i1=5, i2=6, spid=spid, dur=dur_ms)
    return _TIMING_HEAD + "\n" + pars + "\n" + _TIMING_TAIL


def _embed_autoplay_audio(slide, audio_path: Path):
    mime = _MIME_BY_SUFFIX.get(audio_path.suffix.lower(), "audio/mpeg")
    movie = slide.shapes.add_movie(
        str(audio_path),
        Inches(12.7), Inches(0.15), Inches(0.45), Inches(0.45),
        mime_type=mime,
    )
    spid = movie.shape_id
    duration_ms = int(audio_duration_seconds(audio_path) * 1000)
    existing = slide.element.find(qn("p:timing"))
    if existing is not None and _merge_audio_into_timing(existing, spid, duration_ms):
        return  # the deck's own animations stay untouched
    timing = etree.fromstring(_build_timing_xml(spid, duration_ms).encode())
    _insert_slide_child(slide, timing)


def _merge_audio_into_timing(timing, spid: int, dur_ms: int) -> bool:
    """Add narration playback to a slide that already has animations, keeping
    every existing effect. The audio starts when the slide appears, in parallel
    with any auto-start effects; click-triggered effects still wait for clicks.
    Returns False when the timing tree has no shape we recognise."""
    root_ctn = timing.find(f"{qn('p:tnLst')}/{qn('p:par')}/{qn('p:cTn')}")
    if root_ctn is None or root_ctn.get("nodeType") != "tmRoot":
        return False
    root_children = root_ctn.find(qn("p:childTnLst"))
    if root_children is None:
        root_children = etree.SubElement(root_ctn, qn("p:childTnLst"))
    next_id = max((int(c.get("id", 0)) for c in timing.iter(qn("p:cTn"))), default=0) + 1

    main_ctn = next((c for c in root_children.iter(qn("p:cTn"))
                     if c.get("nodeType") == "mainSeq"), None)
    if main_ctn is None:  # only trigger (interactive) sequences: add a main sequence
        seq = etree.fromstring(
            f'<p:seq xmlns:p="{_P}" concurrent="1" nextAc="seek">'
            f'<p:cTn id="{next_id}" dur="indefinite" nodeType="mainSeq"><p:childTnLst/></p:cTn>'
            f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            f'</p:seq>'.encode())
        root_children.insert(0, seq)
        main_ctn = seq.find(qn("p:cTn"))
        next_id += 1
    main_children = main_ctn.find(qn("p:childTnLst"))
    if main_children is None:
        main_children = etree.SubElement(main_ctn, qn("p:childTnLst"))
    main_id = main_ctn.get("id")

    media = etree.fromstring(
        _TIMING_MEDIA.format(i0=next_id, i1=next_id + 1, i2=next_id + 2, spid=spid, dur=dur_ms)
        .replace("<p:par>", f'<p:par xmlns:p="{_P}">', 1).encode())

    # The first step of the main sequence starts on its own when it listens for
    # the sequence's onBegin; join it so audio and auto-start effects run together.
    first = main_children.find(qn("p:par"))
    first_ctn = first.find(qn("p:cTn")) if first is not None else None
    auto_start = first_ctn is not None and any(
        c.get("evt") == "onBegin" for c in first_ctn.iterfind(f"{qn('p:stCondLst')}/{qn('p:cond')}"))
    if auto_start:
        group = first_ctn.find(qn("p:childTnLst"))
        if group is None:
            group = etree.SubElement(first_ctn, qn("p:childTnLst"))
        group.insert(0, media)
        return True
    step = etree.fromstring(
        f'<p:par xmlns:p="{_P}"><p:cTn id="{next_id + 3}" fill="hold">'
        f'<p:stCondLst><p:cond delay="indefinite"/>'
        f'<p:cond evt="onBegin" delay="0"><p:tn val="{main_id}"/></p:cond></p:stCondLst>'
        f'<p:childTnLst/></p:cTn></p:par>'.encode())
    step.find(f"{qn('p:cTn')}/{qn('p:childTnLst')}").append(media)
    main_children.insert(0, step)
    return True


_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _slide_transitions(slide) -> list:
    """Every p:transition on the slide, including the ones PowerPoint wraps in
    mc:AlternateContent (newer transition effects write a p14 choice + fallback)."""
    root = slide.element
    found = list(root.findall(qn("p:transition")))
    for alt in root.findall(f"{{{_MC}}}AlternateContent"):
        found.extend(alt.iter(qn("p:transition")))
    return found


def _set_auto_advance(slide, seconds: float):
    """Advance to the next slide automatically once the narration finishes.
    A transition the deck already has keeps its effect; only the timing is set."""
    adv_ms = int(seconds * 1000)
    existing = _slide_transitions(slide)
    if existing:
        for tr in existing:
            if int(tr.get("advTm", "0") or 0) < adv_ms:
                tr.set("advTm", str(adv_ms))
            tr.set("advClick", tr.get("advClick", "1"))
        return
    transition = etree.fromstring(
        f'<p:transition xmlns:p="{_P}" spd="med" advClick="1" advTm="{adv_ms}"><p:fade/></p:transition>'.encode()
    )
    _insert_slide_child(slide, transition)


# CT_Slide child order: cSld, clrMapOvr, transition, timing, extLst. PowerPoint
# may wrap the transition in mc:AlternateContent, which takes the same slot.
_ORDER = [qn("p:cSld"), qn("p:clrMapOvr"), qn("p:transition"), qn("p:timing"), qn("p:extLst")]
_SLOT = {f"{{{_MC}}}AlternateContent": qn("p:transition")}


def _insert_slide_child(slide, element):
    root = slide.element
    existing = root.find(element.tag)
    if existing is not None:
        root.remove(existing)
    rank = _ORDER.index(element.tag)
    for child in root:
        tag = _SLOT.get(child.tag, child.tag)
        if tag in _ORDER and _ORDER.index(tag) > rank:
            child.addprevious(element)
            return
    root.append(element)
